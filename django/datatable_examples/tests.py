import base64
import datetime
import inspect
import json
from decimal import Decimal
from io import BytesIO
from urllib.parse import urlencode

from ajax_helpers.mixins import AjaxHelpers
from django.contrib.auth.models import AnonymousUser
from django.db.models import Q
from django.db.models.functions import Upper
from django.http import QueryDict
from django.test import RequestFactory, TestCase
from django.urls import reverse
from django_menus.menu import HtmlMenu
from openpyxl import load_workbook

from datatable_examples.column_docs import ALIAS_COLUMNS, COLUMN_SECTIONS, documented_columns
from datatable_examples.models import Company, Payment, Person, Tags
from datatable_examples.views.reference import ColumnReference
from django_datatables import columns as columns_module
from datatable_examples.views.server_side import (JsonBooleanColumn, ServerSideJsonColumn,
                                                  ServerSidePagination, ServerSideTagFilter,
                                                  ServerSideTotalsFilter)
from django_datatables.columns import ManyToManyColumn
from django_datatables.columns import (ColumnLink, DatatableColumn, DateColumn, DateTimeColumn, TickColumn,
                                       LocaleCurrencyColumn, ExcelDatatableColumn, CurrencyColumn,
                                       CurrencyPenceColumn, MultiCurrencyColumn, ZeroPenceColumn, MonthColumn,
                                       YearMonthColumn, AlignColumnLink, ViewLink, SelectColumn,
                                       MultiMenuColumnBase, DatatableColumnError, ChoiceColumn, MenuColumn,
                                       TextFieldColumn, ColumnBase)
from django_datatables.columns import JsonBooleanColumn as LibraryJsonBooleanColumn
from django_datatables.filters import DatatableFilter
from django_datatables.datatables import DatatableTable, DatatableView
from django_datatables.datatables.datatable_error import DatatableError
from django_datatables.filters import PythonPivotFilter
from django_datatables.datatables.server_side import ServerSideTable
from django_datatables.downloads.excel_download import ExcelDownload
from django_datatables.server_side_filters import (ServerDateFilter, ServerPivotFilter, ServerSelect2Filter,
                                                   ServerTagFilter, ServerTotalsFilter, ServerValuesFilter)


def make_data():
    acme = Company.objects.create(name='Acme')
    beta = Company.objects.create(name='Beta')
    blank = Company.objects.create(name='')
    Person.objects.create(company=acme, first_name='Alice', surname='Smith', title=0,
                          date_entered=datetime.date(2020, 1, 10), options={'newsletter': True})
    Person.objects.create(company=acme, first_name='Bob', surname='Jones', title=1,
                          date_entered=datetime.date(2020, 6, 15), options={'newsletter': False})
    Person.objects.create(company=beta, first_name='Carol', surname='Brown', title=0,
                          date_entered=datetime.date(2021, 3, 1))  # no newsletter key
    Person.objects.create(company=blank, first_name='Dave', surname='Green', title=None,
                          date_entered=datetime.date(2021, 12, 31), options={'newsletter': True})
    red = Tags.objects.create(tag='Red')
    blue = Tags.objects.create(tag='Blue')
    red.company.add(acme, beta)
    blue.company.add(acme)
    Payment.objects.create(company=acme, date=datetime.date(2023, 1, 1), amount=100)
    Payment.objects.create(company=acme, date=datetime.date(2023, 2, 1), amount=50)
    Payment.objects.create(company=beta, date=datetime.date(2023, 3, 1), amount=200)
    return acme, beta, blank


def make_table():
    table = ServerSideTable('serverside', model=Person)
    table.add_columns('id', 'first_name', 'surname', 'company__name', 'title_model', 'date_entered')
    return table


def make_edit_table(*edit_fields):
    """A table whose edit_fields make a column added to it editable in place."""
    table = DatatableTable('editing', model=Person)
    table.edit_fields = list(edit_fields) if edit_fields else ['title']
    return table


def make_company_table():
    table = ServerSideTable('companies', model=Company)
    table.add_columns(
        'id',
        'name',
        ManyToManyColumn(column_name='CompanyTags', field='tags__tag', model=Company, html='%1%'),
    )
    return table


def make_custom_tag_table():
    # The 'TagList' column defined in Company.Datatable is a plain DatatableColumn
    # that builds its tag ids in setup_results — it has no ORM field.
    table = ServerSideTable('companies', model=Company)
    table.add_columns('id', 'name', 'TagList')
    return table


def make_json_table():
    # The newsletter column renders options['newsletter'] in row_result, so
    # its value is invisible to the ORM without an explicit field= path.
    table = ServerSideTable('json', model=Person)
    table.add_columns('id', 'first_name', JsonBooleanColumn(column_name='newsletter', json_key='newsletter'))
    return table


class ServerFilterTestCase(TestCase):

    def setUp(self):
        make_data()
        self.table = make_table()

    def get_filter(self, filter_type, column):
        self.table.add_js_filters(filter_type, column)
        return self.table.js_filter_list[-1]

    def names(self, queryset):
        return sorted(r['first_name'] for r in queryset)


class TestFilterSubstitution(ServerFilterTestCase):

    def test_server_classes_substituted(self):
        self.assertIsInstance(self.get_filter('pivot', 'title_model'), ServerPivotFilter)
        self.assertIsInstance(self.get_filter('select2', 'company__name'), ServerSelect2Filter)
        self.assertIsInstance(self.get_filter('date', 'date_entered'), ServerDateFilter)

    def test_tag_and_totals_substituted(self):
        table = make_company_table()
        table.add_js_filters('tag', 'CompanyTags')
        self.assertIsInstance(table.js_filter_list[-1], ServerTagFilter)
        self.table.add_js_filters('totals', 'company__name', sum_column='title_model')
        self.assertIsInstance(self.table.js_filter_list[-1], ServerTotalsFilter)

    def test_unsupported_type_raises(self):
        with self.assertRaises(DatatableError):
            self.table.add_js_filters('expand', 'company__name')

    def test_tag_on_non_m2m_column_skipped(self):
        # The base add_js_filters swallows the DatatableError raised by
        # ServerTagFilter for columns that are not ManyToManyColumns.
        self.table.add_js_filters('tag', 'company__name')
        self.assertEqual(self.table.js_filter_list, [])

    def test_tag_with_explicit_field_on_custom_column(self):
        table = make_custom_tag_table()
        table.add_js_filters('tag', 'TagList', field='tags__pk')
        self.assertIsInstance(table.js_filter_list[-1], ServerTagFilter)

    def test_tag_with_bad_field_path_skipped(self):
        table = make_custom_tag_table()
        table.add_js_filters('tag', 'TagList', field='nonexistent__pk')
        self.assertEqual(table.js_filter_list, [])

    def test_pivot_with_json_key_field(self):
        table = make_json_table()
        table.add_js_filters('pivot', 'newsletter', field='options__newsletter')
        self.assertIsInstance(table.js_filter_list[-1], ServerPivotFilter)

    def test_pivot_with_bad_json_root_skipped(self):
        table = make_json_table()
        table.add_js_filters('pivot', 'newsletter', field='nonexistent__newsletter')
        self.assertEqual(table.js_filter_list, [])

    def test_totals_without_sum_column_skipped(self):
        self.table.add_js_filters('totals', 'company__name')
        self.assertEqual(self.table.js_filter_list, [])

    def test_totals_with_non_orm_sum_column_skipped(self):
        # FullName has a multi-field list as its field so it cannot be summed.
        self.table.add_columns('FullName')
        self.table.add_js_filters('totals', 'company__name', sum_column='FullName')
        self.assertEqual(self.table.js_filter_list, [])

    def test_calculated_column_skipped(self):
        # FullName has a multi-field list as its field; the base add_js_filters
        # swallows the DatatableError so the filter is silently skipped.
        self.table.add_columns('FullName')
        self.table.add_js_filters('pivot', 'FullName')
        self.assertEqual(self.table.js_filter_list, [])


class TestValuesFilter(ServerFilterTestCase):

    def test_in_list_with_choice_labels(self):
        js_filter = self.get_filter('pivot', 'title_model')
        result = js_filter.apply_filter(self.table.get_query(), {'values': ['Mr']})
        self.assertEqual(self.names(result), ['Alice', 'Carol'])

    def test_null_on_int_field(self):
        js_filter = self.get_filter('pivot', 'title_model')
        result = js_filter.apply_filter(self.table.get_query(), {'values': ['null']})
        self.assertEqual(self.names(result), ['Dave'])

    def test_null_on_char_field_includes_empty_string(self):
        js_filter = self.get_filter('select2', 'company__name')
        result = js_filter.apply_filter(self.table.get_query(), {'values': ['null']})
        self.assertEqual(self.names(result), ['Dave'])

    def test_empty_selection_returns_no_rows(self):
        js_filter = self.get_filter('pivot', 'title_model')
        result = js_filter.apply_filter(self.table.get_query(), {'values': []})
        self.assertEqual(result.count(), 0)

    def test_missing_values_is_passthrough(self):
        js_filter = self.get_filter('pivot', 'title_model')
        result = js_filter.apply_filter(self.table.get_query(), {})
        self.assertEqual(result.count(), 4)


class TestFacets(ServerFilterTestCase):

    def test_counts_and_labels(self):
        js_filter = self.get_filter('pivot', 'title_model')
        base = self.table.get_query()
        filtered = base.filter(company__name='Acme')
        facets = js_filter.get_facets(base, filtered)
        self.assertEqual(facets, {'Mr': [1, 2], 'Mrs': [1, 1], 'null': [0, 1]})

    def test_char_facets(self):
        js_filter = self.get_filter('select2', 'company__name')
        base = self.table.get_query()
        facets = js_filter.get_facets(base, base)
        self.assertEqual(facets, {'Acme': [2, 2], 'Beta': [1, 1], 'null': [1, 1]})

    def test_max_facet_values_overflow(self):
        js_filter = self.get_filter('select2', 'company__name')
        js_filter.max_facet_values = 1
        base = self.table.get_query()
        self.assertIsNone(js_filter.get_facets(base, base))


class TestDateFilter(ServerFilterTestCase):

    def test_range_is_strict(self):
        js_filter = self.get_filter('date', 'date_entered')
        base = self.table.get_query()
        result = js_filter.apply_filter(base, {'after': '10/01/2020', 'before': '01/03/2021'})
        # Boundary dates are excluded (matches the client filter's > / <).
        self.assertEqual(self.names(result), ['Bob'])

    def test_malformed_dates_ignored(self):
        js_filter = self.get_filter('date', 'date_entered')
        base = self.table.get_query()
        result = js_filter.apply_filter(base, {'after': 'garbage', 'before': ''})
        self.assertEqual(result.count(), 4)


class TestTagFilter(TestCase):
    """ServerTagFilter on a reverse M2M (Tags.company -> Company).

    Data: Acme has Red + Blue, Beta has Red, the blank-name company is
    untagged.
    """

    def setUp(self):
        make_data()
        self.table = make_company_table()
        self.table.add_js_filters('tag', 'CompanyTags')
        self.js_filter = self.table.js_filter_list[-1]

    def names(self, queryset):
        return sorted(r['name'] for r in queryset)

    def apply(self, **state):
        return self.js_filter.apply_filter(self.table.get_query(), state)

    def test_field_path_reverse_m2m(self):
        self.assertEqual(self.js_filter.field, 'tags__pk')

    def test_include_is_or(self):
        self.assertEqual(self.names(self.apply(include=['Red'])), ['Acme', 'Beta'])

    def test_include_no_duplicate_rows(self):
        # Acme matches both tags; distinct() keeps it to one row.
        result = self.apply(include=['Red', 'Blue'])
        self.assertEqual(self.names(result), ['Acme', 'Beta'])

    def test_required_is_and(self):
        self.assertEqual(self.names(self.apply(required=['Red', 'Blue'])), ['Acme'])

    def test_exclude_drops_any_match(self):
        self.assertEqual(self.names(self.apply(exclude=['Blue'])), ['', 'Beta'])

    def test_include_null_matches_untagged(self):
        self.assertEqual(self.names(self.apply(include=['null'])), [''])

    def test_exclude_null_drops_untagged(self):
        self.assertEqual(self.names(self.apply(exclude=['null'])), ['Acme', 'Beta'])

    def test_empty_state_is_passthrough(self):
        self.assertEqual(self.apply(include=[], required=[], exclude=[]).count(), 3)

    def test_facets_count_distinct_parents(self):
        base = self.table.get_query()
        facets = self.js_filter.get_facets(base, base)
        self.assertEqual(facets, {'Red': [2, 2], 'Blue': [1, 1], 'null': [1, 1]})

    def test_facets_filtered_side(self):
        base = self.table.get_query()
        facets = self.js_filter.get_facets(base, base.filter(name='Acme'))
        self.assertEqual(facets, {'Red': [1, 2], 'Blue': [1, 1], 'null': [0, 1]})

    def test_facets_after_tag_filter_use_fresh_join(self):
        # A queryset filtered on the tag path must still count ALL tags of
        # the matching rows, not just the tags that matched the filter.
        base = self.table.get_query()
        filtered = self.js_filter.apply_filter(base, {'include': ['Red']})
        facets = self.js_filter.get_facets(base, filtered)
        self.assertEqual(facets, {'Red': [2, 2], 'Blue': [1, 1], 'null': [0, 1]})


class TestTagFilterCustomColumn(TestCase):
    """ServerTagFilter with field= on a custom DatatableColumn tag column.

    The 'TagList' column in Company.Datatable builds its tag ids in
    setup_results and carries no ORM field, so the path from Company to the
    tag pks is passed explicitly with field='tags__pk'.
    """

    def setUp(self):
        make_data()
        self.table = make_custom_tag_table()
        self.table.add_js_filters('tag', 'TagList', field='tags__pk')
        self.js_filter = self.table.js_filter_list[-1]

    def names(self, queryset):
        return sorted(r['name'] for r in queryset)

    def apply(self, **state):
        return self.js_filter.apply_filter(self.table.get_query(), state)

    def test_field_is_explicit_path(self):
        self.assertEqual(self.js_filter.field, 'tags__pk')

    def test_include_by_label(self):
        self.assertEqual(self.names(self.apply(include=['Red'])), ['Acme', 'Beta'])

    def test_required_is_and(self):
        self.assertEqual(self.names(self.apply(required=['Red', 'Blue'])), ['Acme'])

    def test_exclude_drops_any_match(self):
        self.assertEqual(self.names(self.apply(exclude=['Blue'])), ['', 'Beta'])

    def test_include_null_matches_untagged(self):
        self.assertEqual(self.names(self.apply(include=['null'])), [''])

    def test_stale_label_matches_nothing(self):
        # A label not in the lookup must not raise from a non-numeric pk query.
        self.assertEqual(self.apply(include=['NoSuchTag']).count(), 0)

    def test_facets_use_labels(self):
        base = self.table.get_query()
        facets = self.js_filter.get_facets(base, base)
        self.assertEqual(facets, {'Red': [2, 2], 'Blue': [1, 1], 'null': [1, 1]})


class TestJsonKeyPivotFilter(TestCase):
    """ServerPivotFilter on a JSON-field key with a choices= label map.

    Data: Alice and Dave have newsletter=true, Bob newsletter=false, Carol
    has no newsletter key at all.  choices= folds False and the missing key
    into the 'No' label, matching the column's row_result.
    """

    def setUp(self):
        make_data()
        self.table = make_json_table()
        self.table.add_js_filters('pivot', 'newsletter', field='options__newsletter',
                                  choices={True: 'Yes', False: 'No', None: 'No'})
        self.js_filter = self.table.js_filter_list[-1]

    def names(self, queryset):
        return sorted(r['first_name'] for r in queryset)

    def apply(self, values):
        return self.js_filter.apply_filter(self.table.get_query(), {'values': values})

    def test_field_is_json_path(self):
        self.assertEqual(self.js_filter.field, 'options__newsletter')

    def test_include_yes(self):
        self.assertEqual(self.names(self.apply(['Yes'])), ['Alice', 'Dave'])

    def test_no_matches_false_and_missing_key(self):
        self.assertEqual(self.names(self.apply(['No'])), ['Bob', 'Carol'])

    def test_yes_and_no_match_all(self):
        self.assertEqual(self.apply(['Yes', 'No']).count(), 4)

    def test_empty_selection_returns_no_rows(self):
        self.assertEqual(self.apply([]).count(), 0)

    def test_unknown_label_matches_nothing(self):
        self.assertEqual(self.apply(['Maybe']).count(), 0)

    def test_facets_fold_missing_key_into_label(self):
        base = self.table.get_query()
        facets = self.js_filter.get_facets(base, base)
        self.assertEqual(facets, {'Yes': [2, 2], 'No': [2, 2]})

    def test_facets_filtered_side(self):
        base = self.table.get_query()
        facets = self.js_filter.get_facets(base, base.filter(company__name='Acme'))
        self.assertEqual(facets, {'Yes': [1, 2], 'No': [1, 2]})


class TestTotalsFilter(TestCase):
    """ServerTotalsFilter summing Person.title (nullable) per company."""

    def setUp(self):
        make_data()
        self.table = make_table()
        self.table.add_js_filters('totals', 'company__name', sum_column='title_model')
        self.js_filter = self.table.js_filter_list[-1]

    def test_facets_are_sums(self):
        base = self.table.get_query()
        facets = self.js_filter.get_facets(base, base)
        # Acme: 0 + 1, Beta: 0, blank company: Sum(None) -> 0.
        self.assertEqual(facets, {'Acme': [1, 1], 'Beta': [0, 0], 'null': [0, 0]})

    def test_apply_filter_matches_pivot(self):
        result = self.js_filter.apply_filter(self.table.get_query(), {'values': ['Acme']})
        self.assertEqual(sorted(r['first_name'] for r in result), ['Alice', 'Bob'])

    def test_max_facet_values_overflow(self):
        self.js_filter.max_facet_values = 1
        base = self.table.get_query()
        self.assertIsNone(self.js_filter.get_facets(base, base))

    def test_facet_value_coercion(self):
        self.assertEqual(self.js_filter._facet_value({'_facet_count': None}), 0)
        self.assertEqual(self.js_filter._facet_value({'_facet_count': Decimal('1.5')}), 1.5)
        self.assertIsInstance(self.js_filter._facet_value({'_facet_count': Decimal('1.5')}), float)
        self.assertEqual(self.js_filter._facet_value({'_facet_count': 7}), 7)


class TestServerSideViewAttribute(TestCase):
    """server_side = True on the view makes the default add_tables create a ServerSideTable."""

    def test_default_add_tables_uses_server_side_table(self):
        view = ServerSidePagination()
        view.add_tables()
        self.assertIsInstance(view.tables['serversidepagination'], ServerSideTable)

    def test_client_side_default_unchanged(self):
        class ClientView(DatatableView):
            model = Person

        view = ClientView()
        view.add_tables()
        table = view.tables['clientview']
        self.assertIsInstance(table, DatatableTable)
        self.assertNotIsInstance(table, ServerSideTable)

    def test_explicit_table_class_overrides_attribute(self):
        class MixedView(DatatableView):
            model = Person
            server_side = True

        view = MixedView()
        view.add_table('client_table', model=Person, table_class=DatatableTable)
        self.assertNotIsInstance(view.tables['client_table'], ServerSideTable)


class TestBaseQuerysetHook(TestCase):
    """base_queryset() lets a table start from an annotated queryset whose annotations then
    behave as plain fields: in the row data, for ordering, and for the server-side filters."""

    def setUp(self):
        make_data()

    def test_default_is_the_model_manager(self):
        table = DatatableTable('people', model=Person)
        self.assertIs(table.base_queryset().model, Person)
        self.assertEqual(table.base_queryset().count(), Person.objects.count())

    def test_annotation_from_base_queryset_is_a_plain_field(self):
        class AnnotatedTable(ServerSideTable):
            def base_queryset(self):
                return Person.objects.annotate(company_upper=Upper('company__name'))

        table = AnnotatedTable('annotated', model=Person)
        table.add_columns('id', 'first_name', 'company_upper')
        rows = list(table.get_query())
        self.assertEqual(len(rows), Person.objects.count())
        for row in rows:
            person = Person.objects.select_related('company').get(pk=row['id'])
            self.assertEqual(row['company_upper'], person.company.name.upper())
        self.assertEqual(table._build_ordering({'order[0][column]': '2', 'order[0][dir]': 'desc'}),
                         ['-company_upper'])
        table.add_js_filters('pivot', 'company_upper')
        pivot = table.js_filter_list[-1]
        self.assertIsInstance(pivot, ServerPivotFilter)
        base = table.get_query()
        filtered = pivot.apply_filter(base, {'values': ['ACME']})
        self.assertEqual({r['company_upper'] for r in filtered}, {'ACME'})
        facets = pivot.get_facets(base, filtered)
        self.assertEqual(facets['ACME'][0], facets['ACME'][1])
        self.assertEqual(sum(v[1] for v in facets.values()), Person.objects.count())


class TestServerSideResponse(TestCase):
    """Simulated DataTables requests through the demo view."""

    def setUp(self):
        make_data()
        self.factory = RequestFactory()

    def post(self, **extra):
        data = {'datatable_data': '1', 'table_id': 'serversidepagination', 'start': '0', 'length': '10',
                'draw': '1'}
        data.update(extra)
        request = self.factory.post('/server-side/pagination', data, HTTP_USER_AGENT='test')
        request.user = AnonymousUser()
        response = ServerSidePagination.as_view()(request)
        return json.loads(response.content)

    def test_first_draw_returns_facets(self):
        result = self.post(draw='1')
        self.assertEqual(result['recordsTotal'], 4)
        self.assertIn('company__name', result['facets'])
        self.assertIn('title_model', result['facets'])
        self.assertNotIn('date_entered', result['facets'])

    def test_later_draw_skips_facets_unless_requested(self):
        self.assertNotIn('facets', self.post(draw='2'))
        self.assertIn('facets', self.post(draw='2', need_facets='1'))

    def test_filter_state_changes_records_filtered(self):
        state = json.dumps({'title_model': {'type': 'pivot', 'values': ['Mr']}})
        result = self.post(js_filter_state=state)
        self.assertEqual(result['recordsFiltered'], 2)
        self.assertEqual(result['recordsTotal'], 4)
        self.assertEqual(result['facets']['company__name'], {'Acme': [1, 2], 'Beta': [1, 1], 'null': [0, 1]})

    def test_search_and_filter_compose(self):
        state = json.dumps({'title_model': {'type': 'pivot', 'values': ['Mr']}})
        result = self.post(js_filter_state=state, **{'search[value]': 'Carol'})
        self.assertEqual(result['recordsFiltered'], 1)

    def test_invalid_filter_state_ignored(self):
        result = self.post(js_filter_state='not json')
        self.assertEqual(result['recordsFiltered'], 4)

    def test_date_filter_applied(self):
        state = json.dumps({'date_entered': {'type': 'date', 'after': '01/01/2021', 'before': ''}})
        result = self.post(js_filter_state=state)
        self.assertEqual(result['recordsFiltered'], 2)

    def test_length_all_returns_every_row(self):
        # DataTables sends length=-1 for an "All" lengthMenu entry; a negative
        # slice bound would raise, so it must be treated as the page cap.
        result = self.post(length='-1')
        self.assertEqual(len(result['data']), 4)


class ServerSideViewTestCase(TestCase):
    view_class = None
    table_id = None
    url = '/'

    def setUp(self):
        make_data()
        self.factory = RequestFactory()

    def post(self, **extra):
        data = {'datatable_data': '1', 'table_id': self.table_id, 'start': '0', 'length': '10', 'draw': '1'}
        data.update(extra)
        request = self.factory.post(self.url, data, HTTP_USER_AGENT='test')
        request.user = AnonymousUser()
        response = self.view_class.as_view()(request)
        return json.loads(response.content)


class TestServerSideTagResponse(ServerSideViewTestCase):
    view_class = ServerSideTagFilter
    table_id = 'serversidetagfilter'
    url = '/server-side/tag-filter'

    def test_first_draw_returns_tag_facets(self):
        result = self.post(draw='1')
        self.assertEqual(result['recordsTotal'], 3)
        self.assertEqual(result['facets']['CompanyTags'], {'Red': [2, 2], 'Blue': [1, 1], 'null': [1, 1]})

    def test_tag_filter_state_applied(self):
        state = json.dumps({'CompanyTags': {'type': 'tag', 'include': ['Red'], 'required': [], 'exclude': []}})
        result = self.post(js_filter_state=state)
        self.assertEqual(result['recordsFiltered'], 2)
        self.assertEqual(result['recordsTotal'], 3)
        self.assertEqual(result['facets']['CompanyTags'], {'Red': [2, 2], 'Blue': [1, 1], 'null': [0, 1]})

    def test_tag_filter_no_duplicate_rows(self):
        state = json.dumps({'CompanyTags': {'type': 'tag', 'include': ['Red', 'Blue'],
                                            'required': [], 'exclude': []}})
        result = self.post(js_filter_state=state)
        self.assertEqual(result['recordsFiltered'], 2)
        self.assertEqual(len(result['data']), 2)


class TestServerSideTotalsResponse(ServerSideViewTestCase):
    view_class = ServerSideTotalsFilter
    table_id = 'serversidetotalsfilter'
    url = '/server-side/totals-filter'

    def test_first_draw_returns_sum_facets(self):
        result = self.post(draw='1')
        self.assertEqual(result['recordsTotal'], 3)
        self.assertEqual(result['facets']['company__name'], {'Acme': [150, 150], 'Beta': [200, 200]})

    def test_totals_filter_state_applied(self):
        state = json.dumps({'company__name': {'type': 'totals', 'values': ['Acme']}})
        result = self.post(js_filter_state=state)
        self.assertEqual(result['recordsFiltered'], 2)
        self.assertEqual(result['facets']['company__name'], {'Acme': [150, 150], 'Beta': [0, 200]})


class TestServerSideJsonResponse(ServerSideViewTestCase):
    view_class = ServerSideJsonColumn
    table_id = 'serversidejsoncolumn'
    url = '/server-side/json-column'

    def test_first_draw_returns_json_key_facets(self):
        result = self.post(draw='1')
        self.assertEqual(result['recordsTotal'], 4)
        self.assertEqual(result['facets']['newsletter'], {'Yes': [2, 2], 'No': [2, 2]})

    def test_json_filter_state_applied(self):
        state = json.dumps({'newsletter': {'type': 'pivot', 'values': ['No']}})
        result = self.post(js_filter_state=state)
        self.assertEqual(result['recordsFiltered'], 2)
        self.assertEqual(result['facets']['newsletter'], {'Yes': [0, 2], 'No': [2, 2]})

    def test_rendered_rows_match_filter(self):
        state = json.dumps({'newsletter': {'type': 'pivot', 'values': ['Yes']}})
        result = self.post(js_filter_state=state)
        self.assertEqual(sorted(row[1] for row in result['data']), ['Alice', 'Dave'])


class TestFilteredQuery(TestCase):
    """ServerSideTable.filtered_query rebuilds the queryset from a DataTables request state."""

    def setUp(self):
        make_data()
        self.table = make_table()

    def query(self, **state):
        return self.table.filtered_query(QueryDict(urlencode(state)))

    def test_no_state_returns_everything(self):
        self.assertEqual(self.query().count(), 4)

    def test_search_and_js_filter_compose(self):
        state = json.dumps({'title_model': {'type': 'pivot', 'values': ['Mr']}})
        self.table.add_js_filters('pivot', 'title_model')
        self.table.search_fields = ['first_name', 'surname', 'company__name']
        results = self.query(js_filter_state=state, **{'search[value]': 'Carol'})
        self.assertEqual([r['first_name'] for r in results], ['Carol'])

    def test_column_search_and_ordering(self):
        results = self.query(**{'columns[0][name]': 'company__name', 'columns[0][search][value]': 'acme',
                                'order[0][column]': '2', 'order[0][dir]': 'desc'})
        self.assertEqual([r['surname'] for r in results], ['Smith', 'Jones'])


def make_link_table(**link_kwargs):
    """Server-side table with a list-field ColumnLink (id -> URL, surname -> display).
    person_link is column index 2."""
    table = ServerSideTable('serverside', model=Person)
    table.add_columns(
        'id',
        'first_name',
        ColumnLink(column_name='person_link', field=['id', 'surname'],
                   url_name='column_visibility', **link_kwargs),
        'company__name',
    )
    return table


class TestSearchField(TestCase):
    """search_field / search / get_search_filter make complex columns (list-field links,
    calculated fields, custom Q) searchable and sortable server-side."""

    def setUp(self):
        make_data()

    def query(self, table, **state):
        return list(table.filtered_query(QueryDict(urlencode(state))))

    # ---- _search_paths / get_search_filter resolution ----

    def test_plain_field_is_searchable(self):
        table = make_link_table()
        col, _ = table.find_column('first_name')
        self.assertEqual(col._search_paths(), ['first_name'])
        self.assertIsNotNone(col.get_search_filter('a'))

    def test_calculated_column_not_searchable(self):
        col = DatatableColumn(column_name='_total', field='amount', model=Person)
        self.assertTrue(col.options.get('calculated'))
        self.assertEqual(col._search_paths(), [])
        self.assertIsNone(col.get_search_filter('x'))

    def test_list_field_columnlink_defaults_to_display_field(self):
        col, _ = make_link_table().find_column('person_link')
        self.assertEqual(col._search_paths(), ['surname'])

    def test_explicit_search_field_overrides_default(self):
        col, _ = make_link_table(search_field='first_name').find_column('person_link')
        self.assertEqual(col._search_paths(), ['first_name'])

    def test_search_field_false_disables(self):
        col, _ = make_link_table(search_field=False).find_column('person_link')
        self.assertEqual(col._search_paths(), [])
        self.assertIsNone(col.get_search_filter('smith'))

    def test_search_field_list_ors_icontains(self):
        col = DatatableColumn(column_name='name', field='surname',
                              search_field=['first_name', 'surname'], model=Person)
        self.assertEqual(col._search_paths(), ['first_name', 'surname'])
        # 'brown' matches surname Brown; 'bob' matches first_name Bob — proves the OR.
        self.assertEqual([p.first_name for p in Person.objects.filter(col.get_search_filter('brown'))],
                         ['Carol'])
        self.assertEqual([p.first_name for p in Person.objects.filter(col.get_search_filter('bob'))],
                         ['Bob'])

    def test_search_callable_returns_q(self):
        col = DatatableColumn(column_name='name', field='surname', model=Person,
                              search=lambda v: Q(first_name__icontains=v) | Q(surname__icontains=v))
        self.assertEqual([p.first_name for p in Person.objects.filter(col.get_search_filter('brown'))],
                         ['Carol'])

    def test_search_callable_dict_wrapped_in_q(self):
        col = DatatableColumn(column_name='name', field='surname', model=Person,
                              search=lambda v: {'first_name__iexact': v})
        q = col.get_search_filter('Bob')
        self.assertIsInstance(q, Q)
        self.assertEqual([p.first_name for p in Person.objects.filter(q)], ['Bob'])

    def test_relation_column_field_not_double_prefixed(self):
        col, _ = make_link_table().find_column('company__name')
        self.assertEqual(col._search_paths(), ['company__name'])

    # ---- search-box visibility (no_col_search auto-marking) ----

    def test_columnlink_search_box_shown(self):
        col, _ = make_link_table().find_column('person_link')
        self.assertNotEqual(col.options.get('no_col_search'), True)

    def test_search_field_false_hides_box(self):
        col, _ = make_link_table(search_field=False).find_column('person_link')
        self.assertTrue(col.options.get('no_col_search'))

    def test_callable_only_column_box_shown_but_not_sortable(self):
        # search_field=False removes the sortable path; the search= callable still
        # makes the column searchable (box shown).
        table = ServerSideTable('t', model=Person)
        table.add_columns(
            'id',
            DatatableColumn(column_name='name_search', field='surname', search_field=False,
                            search=lambda v: Q(surname__icontains=v)),
        )
        col, _ = table.find_column('name_search')
        self.assertNotEqual(col.options.get('no_col_search'), True)
        self.assertEqual(col._search_paths(), [])

    # ---- filtered_query end to end ----

    def test_columnlink_header_search_filters_display_field(self):
        results = self.query(make_link_table(),
                             **{'columns[0][name]': 'person_link', 'columns[0][search][value]': 'smith'})
        self.assertEqual([r['surname'] for r in results], ['Smith'])

    def test_explicit_search_field_used_in_header_search(self):
        results = self.query(make_link_table(search_field='first_name'),
                             **{'columns[0][name]': 'person_link', 'columns[0][search][value]': 'carol'})
        self.assertEqual([r['first_name'] for r in results], ['Carol'])

    def test_columnlink_ordering_sorts_by_display_field(self):
        results = self.query(make_link_table(), **{'order[0][column]': '2', 'order[0][dir]': 'desc'})
        self.assertEqual([r['surname'] for r in results], ['Smith', 'Jones', 'Green', 'Brown'])

    def test_search_callable_applied_in_header_search(self):
        table = ServerSideTable('t', model=Person)
        table.add_columns(
            'id',
            DatatableColumn(column_name='name_search', field='surname',
                            search=lambda v: Q(first_name__icontains=v) | Q(surname__icontains=v)),
        )
        # 'jones' hits the surname OR-branch (Bob Jones); no first_name matches.
        results = self.query(table,
                             **{'columns[0][name]': 'name_search', 'columns[0][search][value]': 'jones'})
        self.assertEqual([r['surname'] for r in results], ['Jones'])

    def test_global_search_includes_columnlink_display_field(self):
        # No search_fields set -> each column's get_search_filter builds the predicate.
        results = self.query(make_link_table(), **{'search[value]': 'brown'})
        self.assertEqual([r['surname'] for r in results], ['Brown'])


class TestPythonPivotFilter(TestCase):
    """PythonPivotFilter renders its fixed checkbox list from package templates."""

    @staticmethod
    def make_filter(filter_list):
        table = DatatableTable('people', model=Person)
        table.add_columns('id', 'first_name')
        table.add_js_filters(PythonPivotFilter, 'first_name', filter_list=filter_list)
        return table.js_filter_list[-1]

    def test_renders_checkboxes_and_options(self):
        html = self.make_filter([('Active', 'active'), ('Other', 'other')]).render()
        self.assertIn('data-value="active"', html)
        self.assertIn('checked>Active', html)
        self.assertIn('["active", "other"]', html)
        self.assertIn('PythonPivotFilter', html)

    def test_values_are_url_encoded(self):
        html = self.make_filter([('New York', 'New York')]).render()
        self.assertIn('data-value="New%20York"', html)

    def test_option_json_cannot_close_script_tag(self):
        html = self.make_filter([('Bad', '</script>')]).render()
        self.assertNotIn('["</script>"]', html)
        self.assertIn('\\u003C/script>', html)

    def test_server_side_table_rejects_client_side_filter_class(self):
        table = ServerSideTable('people', model=Person)
        table.add_columns('id', 'first_name')
        with self.assertRaises(DatatableError):
            table.add_js_filters(PythonPivotFilter, 'first_name', filter_list=[('A', 'a')])


class TestServerSideExcelDownload(TestCase):
    """The download button posts the DataTables request state; the export covers the whole filtered dataset."""

    def setUp(self):
        make_data()
        self.factory = RequestFactory()

    def download(self, **state):
        data = {'column': 'get_excel', 'table_id': 'serversidepagination',
                'column_data': '[]', 'datatable_state': urlencode(state)}
        request = self.factory.post('/server-side/pagination', json.dumps(data), content_type='application/json',
                                    HTTP_X_REQUESTED_WITH='XMLHttpRequest', HTTP_USER_AGENT='test')
        request.user = AnonymousUser()
        response = ServerSidePagination.as_view()(request)
        commands = json.loads(response.content)
        save_file = next(c for c in commands if c['function'] == 'save_file')
        workbook = load_workbook(BytesIO(base64.b64decode(save_file['data'])))
        return list(workbook.active.values)

    def test_download_includes_all_rows_not_just_current_page(self):
        rows = self.download(start='0', length='2')
        self.assertEqual(len(rows), 5)  # header + all 4 people despite a 2-row page

    def test_download_applies_js_filters(self):
        state = json.dumps({'title_model': {'type': 'pivot', 'values': ['Mr']}})
        rows = self.download(js_filter_state=state)
        self.assertEqual(len(rows), 3)

    def test_download_applies_global_search(self):
        rows = self.download(**{'search[value]': 'Carol'})
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][2], 'Brown')

    def test_download_applies_ordering(self):
        rows = self.download(**{'order[0][column]': '2', 'order[0][dir]': 'desc'})
        self.assertEqual([r[2] for r in rows[1:]], ['Smith', 'Jones', 'Green', 'Brown'])

    def test_link_column_exports_display_text(self):
        # person_link (column 3) is a list-field ColumnLink whose row_result is [id, surname];
        # the export must be the surname, not the repr of the pair.
        rows = self.download(**{'search[value]': 'Carol'})
        self.assertEqual(rows[1][3], 'Brown')


class ExcelStyleView(ExcelDownload, AjaxHelpers, DatatableView):
    """A hidden column in front of two styled columns."""

    model = Payment
    ajax_commands = ['column']

    @staticmethod
    def setup_table(table):
        table.add_columns(
            '.id',  # hidden, so dropped from the export
            'company__name',
            CurrencyColumn(column_name='amount', field='amount'),
            DateColumn(column_name='date', field='date'),
        )


class TestExcelStyles(TestCase):
    """xl_style must land on the cell its column actually wrote, which is not its index in
    table.columns once a hidden column has been dropped from the row."""

    def setUp(self):
        make_data()
        self.factory = RequestFactory()

    def download(self):
        data = {'column': 'get_excel', 'table_id': 'excelstyleview',
                'column_data': json.dumps(list(Payment.objects.values_list('id', flat=True)))}
        request = self.factory.post('/', json.dumps(data), content_type='application/json',
                                    HTTP_X_REQUESTED_WITH='XMLHttpRequest', HTTP_USER_AGENT='test')
        request.user = AnonymousUser()
        response = ExcelStyleView.as_view()(request)
        save_file = next(c for c in json.loads(response.content) if c['function'] == 'save_file')
        return load_workbook(BytesIO(base64.b64decode(save_file['data']))).active

    def test_hidden_column_is_not_exported(self):
        sheet = self.download()
        self.assertEqual([c.value for c in sheet[1]], ['Name', 'Amount', 'Date'])

    def test_styles_follow_their_own_column(self):
        sheet = self.download()
        for row in sheet.iter_rows(min_row=2):
            self.assertEqual(row[1].number_format, '#,##0.00_-')
            self.assertEqual(row[2].number_format, 'dd/mm/yyyy')

    def test_currency_and_date_export_as_values_not_text(self):
        sheet = self.download()
        for row in sheet.iter_rows(min_row=2):
            self.assertIsInstance(row[1].value, (int, float))  # a number, not '100.00'
            self.assertIsInstance(row[2].value, datetime.date)


class TestPortedColumns(TestCase):
    """Unit tests for the columns / helpers consolidated into the library."""

    def setUp(self):
        make_data()

    def test_date_column_render_and_excel(self):
        col = DateColumn('date_entered', column_name='Date', model=Person)
        self.assertEqual(col.row_result({'date_entered': datetime.date(2020, 1, 10)}, {}), '10/01/2020')
        self.assertEqual(col.excel('10/01/2020'), datetime.date(2020, 1, 10))
        self.assertEqual(col.excel(''), '')

    def test_date_column_format_hook_overridable(self):
        class UsDate(DateColumn):
            date_str = '%m/%d/%Y'
        col = UsDate('date_entered', column_name='Date', model=Person)
        self.assertEqual(col.row_result({'date_entered': datetime.date(2020, 1, 10)}, {}), '01/10/2020')

    def test_datetime_column_render_and_excel(self):
        col = DateTimeColumn('date_entered', column_name='DT', model=Person)
        dt = datetime.datetime(2020, 1, 10, 14, 30)
        self.assertEqual(col.row_result({'date_entered': dt}, {}), '10/01/2020 14:30')
        self.assertEqual(col.excel('10/01/2020 14:30'), dt)

    def test_tick_column(self):
        col = TickColumn(column_name='active', field='id', model=Person)
        self.assertEqual(col.row_result({'id': 1}, {}), 'true')
        self.assertEqual(col.row_result({'id': 0}, {}), 'false')
        self.assertIn('fa-check-circle', dict(col.options['lookup'])['true'])

    def test_json_boolean_column(self):
        col = LibraryJsonBooleanColumn(column_name='newsletter', json_key='newsletter', field='options', model=Person)
        self.assertEqual(col.row_result({'options': {'newsletter': True}}, {}), 'Yes')
        self.assertEqual(col.row_result({'options': {'newsletter': False}}, {}), 'No')
        self.assertEqual(col.row_result({'options': {}}, {}), 'No')

    def test_locale_currency_column(self):
        col = LocaleCurrencyColumn(column_name='amount', field='amount', model=Payment)
        self.assertEqual(col.row_result({'amount': 12345}, {}), '123.4500')
        render = col.options['render'][0]
        self.assertEqual(render['currency_code'], 'GBP')
        self.assertEqual(render['locale'], 'en_GB')
        self.assertEqual(col.excel('123.45'), 123.45)

    def test_excel_datatable_column_strips_html(self):
        col = ExcelDatatableColumn(column_name='x', field='name', model=Company)
        self.assertEqual(col.excel('<b>Hello</b><br><i>World</i>'), 'Hello, World')
        self.assertEqual(col.excel(None), '')

    def test_column_link_new_tab(self):
        col = ColumnLink(column_name='name', field='name', url_name='column_visibility', new_tab=True, model=Company)
        self.assertIn('target="_blank"', col.options['render'][0]['html'])
        plain = ColumnLink(column_name='name', field='name', url_name='column_visibility', model=Company)
        self.assertNotIn('target="_blank"', plain.options['render'][0]['html'])

    def test_many_to_many_refresh_picks_up_new_related_rows(self):
        col = ManyToManyColumn(column_name='CompanyTags', field='tags__tag', model=Company)
        self.assertNotIn('Green', col.options['lookup_dict'].values())
        green = Tags.objects.create(tag='Green')
        green.company.add(Company.objects.first())
        col.refresh_lookup()
        self.assertIn('Green', col.options['lookup_dict'].values())

    def test_many_to_many_excel(self):
        col = ManyToManyColumn(column_name='CompanyTags', field='tags__tag', model=Company)
        red = Tags.objects.get(tag='Red')
        self.assertEqual(col.excel([red.pk]), 'Red')

    def test_many_to_many_excel_exports_the_blank_text(self):
        # The blank entry only ever reaches options['lookup'], never lookup_dict.
        col = ManyToManyColumn(column_name='CompanyTags', field='tags__tag', model=Company, blank='No tags')
        self.assertEqual(col.excel(col.blank), 'No tags')

    def test_filter_column_titles(self):
        table = make_table()
        f = DatatableFilter('pivot', table, columns=['surname', 'first_name'])
        self.assertEqual([t[1] for t in f.column_titles()], ['surname', 'first_name'])


class TestColumnExcelValues(TestCase):
    """Every column's excel() must turn its own row_result into something Excel can hold -
    no list reprs, no markup, no numbers stranded as text."""

    def setUp(self):
        make_data()

    def test_link_column_exports_display_text_not_the_ref(self):
        col, _ = make_link_table().find_column('person_link')
        self.assertEqual(col.row_result({'id': 7, 'surname': 'Smith'}, {}), [7, 'Smith'])
        self.assertEqual(col.excel([7, 'Smith']), 'Smith')

    def test_link_column_with_a_scalar_field_is_unaffected(self):
        col = ColumnLink(column_name='name', field='name', url_name='column_visibility', model=Company)
        self.assertEqual(col.excel('Acme'), 'Acme')

    def test_choice_column_edit_exports_the_label(self):
        col = ChoiceColumn(column_name='title', field='title', model=Person,
                           choices=[(0, 'Mr'), (1, 'Mrs')], table=make_edit_table())
        self.assertEqual(col.row_result({'title': 0}, {}), [0, 'Mr'])
        self.assertEqual(col.excel([0, 'Mr']), 'Mr')

    def test_editable_fk_column_exports_the_label(self):
        # setup_edit rewrites the field to [id, label], so row_result becomes a pair.
        table = make_edit_table('company__name')
        table.add_columns('company__name')
        col, _ = table.find_column('company__name')
        self.assertEqual(col.edit_type, 'FK')
        self.assertEqual(col.excel([3, 'Acme']), 'Acme')

    def test_currency_column_exports_a_number(self):
        col = CurrencyPenceColumn(column_name='amount', field='amount', model=Payment)
        self.assertEqual(col.excel(col.row_result({'amount': 12345}, {})), 123.45)
        self.assertEqual(col.excel(None), '')

    def test_locale_currency_column_with_a_list_field(self):
        # row_result returns (amount, currency) here; float() used to choke on the pair.
        col = LocaleCurrencyColumn(column_name='amount', field=['amount', 'currency'], model=Payment)
        self.assertEqual(col.excel(col.row_result({'amount': 12345, 'currency': 'USD'}, {})), 123.45)

    def test_text_field_column_exports_newlines_not_br_tags(self):
        col = TextFieldColumn(column_name='notes', field='name', model=Company)
        self.assertEqual(col.row_result({'name': 'one\ntwo'}, {}), 'one<br>two')
        self.assertEqual(col.excel('one<br>two'), 'one\ntwo')
        self.assertEqual(col.excel(None), '')

    def test_button_columns_are_not_exported(self):
        select = SelectColumn(column_name='select', model=Company)
        self.assertTrue(select.xl_dont_show())
        menu = MenuColumn(column_name='menu', field='id', model=Company,
                          menu=HtmlMenu(RequestFactory().get('/'), 'button_group').add_items(
                              ('column_visibility', 'View')))
        self.assertTrue(menu.xl_dont_show())

    def test_ordinary_column_is_exported(self):
        self.assertFalse(DatatableColumn(column_name='name', field='name', model=Company).xl_dont_show())


class TestColumnReference(TestCase):
    """The Column Reference page is hand written, so keep it in step with the package."""

    @staticmethod
    def exported_columns():
        # ColumnNameError / DatatableColumnError are Exceptions and EDIT_CELL_HTML is a str, so the
        # ColumnBase check leaves exactly the column classes.
        return {obj for _, obj in inspect.getmembers(columns_module, inspect.isclass)
                if issubclass(obj, ColumnBase)}

    def test_every_exported_column_is_documented(self):
        missing = self.exported_columns() - set(documented_columns()) - set(ALIAS_COLUMNS)
        self.assertEqual(
            missing, set(),
            f'{sorted(c.__name__ for c in missing)} is exported from django_datatables.columns but is not on '
            f'the Column Reference page. Add a ColumnDoc to the right Section in column_docs.py, or add it to '
            f'ALIAS_COLUMNS if it is an alias of a column that is already documented.')

    def test_modal_columns_are_documented(self):
        # Not exported from django_datatables.columns, so the sweep above cannot police them.
        self.assertLessEqual({'ModalLink', 'ColumnLinkCoalesce'}, {c.__name__ for c in documented_columns()})

    def test_aliases_are_not_also_documented(self):
        both = set(ALIAS_COLUMNS) & set(documented_columns())
        self.assertEqual(both, set(), f'{sorted(c.__name__ for c in both)} is in both ALIAS_COLUMNS and a '
                                      f'Section - it belongs in one or the other.')

    def test_no_column_is_documented_twice(self):
        documented = documented_columns()
        self.assertEqual(len(documented), len(set(documented)), 'a column has two ColumnDoc entries')

    def test_entries_are_columns_with_prose(self):
        for section in COLUMN_SECTIONS:
            for doc in section.columns:
                self.assertTrue(issubclass(doc.column, ColumnBase), f'{doc.name} is not a column class')
                self.assertTrue(doc.summary, f'{doc.name} has no summary')
                self.assertTrue(doc.example, f'{doc.name} has no example')

    def test_demo_links_resolve(self):
        # A renamed url_name should fail here rather than 500 the page.
        for section in COLUMN_SECTIONS:
            for doc in section.columns:
                if doc.demo:
                    reverse(doc.demo)
                    self.assertTrue(doc.demo_title, f'{doc.name} links to {doc.demo} with no title')

    def test_page_renders_every_column(self):
        request = RequestFactory().get('/reference/columns', HTTP_USER_AGENT='test')
        request.user = AnonymousUser()
        response = ColumnReference.as_view()(request)
        self.assertEqual(response.status_code, 200)
        # as_view() returns an unrendered TemplateResponse.
        content = response.render().content.decode()
        for column in documented_columns():
            self.assertIn(column.__name__, content)


class TestMergedColumns(TestCase):
    """Behaviour of the currency / boolean / link column families after de-duplication."""

    def test_currency_columns_share_a_divisor(self):
        pence = CurrencyPenceColumn(column_name='amount', field='amount', model=Payment)
        units = CurrencyColumn(column_name='amount', field='amount', model=Payment)
        self.assertEqual(pence.row_result({'amount': 12345}, {}), '123.45')
        self.assertEqual(units.row_result({'amount': 123.45}, {}), '123.45')
        for col in (pence, units):
            self.assertEqual(col.column_defs['className'], 'dt-right')
            self.assertIsNone(col.row_result({'amount': None}, {}))
            self.assertIsNone(col.row_result({}, {}))

    def test_currency_column_keeps_width_kwarg(self):
        # setup_kwargs used to replace column_defs wholesale, discarding width.
        col = CurrencyPenceColumn(column_name='amount', field='amount', width='80px', model=Payment)
        self.assertEqual(col.column_defs['width'], '80px')
        self.assertEqual(col.column_defs['className'], 'dt-right')

    def test_zero_pence_column_formats_missing_value(self):
        col = ZeroPenceColumn(column_name='amount', field='amount', model=Payment)
        self.assertEqual(col.row_result({'amount': 12345}, {}), '123.45')
        self.assertEqual(col.row_result({'amount': None}, {}), '0.00')

    def test_multi_currency_column_honours_pennies(self):
        # pennies=True means the value is already in whole units, so it must not be divided.
        col = MultiCurrencyColumn(column_name='amount', field=['amount', 'currency'], model=Payment)
        self.assertEqual(col.row_result({'amount': 12345, 'currency': 'USD'}, {}), (123.45, 'USD'))
        whole = MultiCurrencyColumn(column_name='amount', field=['amount', 'currency'], pennies=True, model=Payment)
        self.assertEqual(whole.row_result({'amount': 123, 'currency': 'USD'}, {}), (123, 'USD'))

    def test_multi_currency_column_render_matches_locale_column(self):
        col = MultiCurrencyColumn(column_name='amount', field=['amount', 'currency'], model=Payment)
        render = col.options['render'][0]
        self.assertTrue(render['multi_currency'])
        self.assertEqual(render['function'], 'currency')
        self.assertEqual(render['currency_code'], 'GBP')
        self.assertEqual(render['locale'], 'en_GB')
        self.assertNotIn('multi_currency', LocaleCurrencyColumn(column_name='amount', field='amount',
                                                                model=Payment).options['render'][0])

    def test_month_and_year_month_columns_agree(self):
        month = MonthColumn(column_name='m', field='date_entered', model=Person)
        year_month = YearMonthColumn(column_name='ym', field='date_entered', model=Person)
        data = {'date_entered': datetime.date(2020, 1, 10)}
        self.assertEqual(month.row_result(data, {}), '2020 01')
        self.assertEqual(year_month.row_result(data, {}), month.row_result(data, {}))
        self.assertTrue(year_month.options['hidden'])
        self.assertNotIn('hidden', month.options)
        self.assertEqual(year_month.row_result({'date_entered': None}, {}), '')

    def test_column_link_align_kwarg(self):
        plain = ColumnLink(column_name='name', field='name', url_name='column_visibility', model=Company)
        self.assertNotIn('className', plain.column_defs)
        aligned = ColumnLink(column_name='name', field='name', url_name='column_visibility', align='center',
                             model=Company)
        self.assertEqual(aligned.column_defs['className'], 'dt-center')

    def test_align_column_link_and_view_link_defaults(self):
        centre = AlignColumnLink(column_name='name', field='name', url_name='column_visibility', model=Company)
        self.assertEqual(centre.column_defs['className'], 'dt-center')
        view = ViewLink(column_name='name', url_name='column_visibility', model=Company)
        self.assertEqual(view.column_defs['className'], 'dt-right')
        self.assertIn('View', view.options['render'][0]['html'])
        # An explicit align= still wins over the class default.
        left = ViewLink(column_name='name', url_name='column_visibility', align='left', model=Company)
        self.assertEqual(left.column_defs['className'], 'dt-left')

    def test_view_link_has_no_dead_css_class_kwarg(self):
        view = ViewLink(column_name='name', url_name='column_visibility', model=Company)
        self.assertNotIn('css_class', view.options)

    def test_tick_column_replace_overridable(self):
        col = TickColumn(column_name='active', field='id', replace=('Y', 'N'), model=Person)
        self.assertEqual(dict(col.options['lookup'])['true'], 'Y')

    def test_json_boolean_column_custom_choices(self):
        col = LibraryJsonBooleanColumn(column_name='newsletter', json_key='newsletter', field='options',
                                       choices=['On', 'Off'], model=Person)
        self.assertEqual(col.row_result({'options': {'newsletter': True}}, {}), 'On')
        self.assertEqual(col.row_result({'options': {}}, {}), 'Off')
        self.assertEqual(col.column_defs['width'], '60px')
        self.assertTrue(col.options['no_col_search'])

    def test_multi_menu_column_base_is_a_no_heading_column(self):
        col = MultiMenuColumnBase(column_name='menu', field='id', model=Person)
        self.assertEqual(col.title, '')
        self.assertEqual(col.column_defs['width'], '120px')
        self.assertFalse(col.column_defs['orderable'])
        self.assertTrue(col.options['no_col_search'])
        self.assertEqual(col.options['lookup'], [])

    def test_multi_menu_column_base_requires_a_view_for_menus(self):
        col = MultiMenuColumnBase(column_name='menu', field='id', model=Person)
        with self.assertRaises(DatatableColumnError):
            col.add_menu(1, ('column_visibility', 'View'))

    def test_json_boolean_column_handles_null_json_field(self):
        col = LibraryJsonBooleanColumn(column_name='newsletter', json_key='newsletter', field='options', model=Person)
        self.assertEqual(col.row_result({'options': None}, {}), 'No')

    def test_demo_json_boolean_column_is_the_library_one(self):
        # The demo used to carry its own near-identical copy.
        self.assertIs(JsonBooleanColumn, LibraryJsonBooleanColumn)

    def test_select_column_title_overridable(self):
        self.assertIn('Select all', SelectColumn(column_name='sel', model=Person).title)
        self.assertEqual(SelectColumn(column_name='sel', title='', model=Person).title, '')
